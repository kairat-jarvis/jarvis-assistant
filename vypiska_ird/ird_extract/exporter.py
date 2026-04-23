"""Экспорт результата в DOCX и XLSX."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Dict

from docx import Document
from docx.shared import Cm, Pt
from openpyxl import Workbook

from .fields import FIELDS


def export_docx(data: Dict[str, str], output_path: str, source_name: str = "") -> None:
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    heading = doc.add_heading("Выписка из ИРД", level=1)
    heading.alignment = 1  # center

    meta = doc.add_paragraph()
    meta.add_run(f"Источник: {source_name or '—'}\n").italic = True
    meta.add_run(f"Дата формирования: {datetime.now().strftime('%d.%m.%Y %H:%M')}").italic = True

    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    table.columns[0].width = Cm(7)
    table.columns[1].width = Cm(10)

    hdr = table.rows[0].cells
    hdr[0].text = "Поле"
    hdr[1].text = "Значение"

    for field in FIELDS:
        row = table.add_row().cells
        row[0].text = field["label"]
        row[1].text = data.get(field["key"], "") or "—"

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def export_xlsx(data: Dict[str, str], output_path: str, source_name: str = "") -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Выписка ИРД"

    ws["A1"] = "Источник"
    ws["B1"] = source_name or ""
    ws["A2"] = "Дата"
    ws["B2"] = datetime.now().strftime("%d.%m.%Y %H:%M")

    ws["A4"] = "Поле"
    ws["B4"] = "Значение"
    for idx, field in enumerate(FIELDS, start=5):
        ws.cell(row=idx, column=1, value=field["label"])
        ws.cell(row=idx, column=2, value=data.get(field["key"], ""))

    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 60

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
